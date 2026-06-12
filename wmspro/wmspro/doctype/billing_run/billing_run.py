# # Copyright (c) 2026, Quantbit Technologies Private Limited  and contributors
# # For license information, please see license.txt

import frappe
from frappe.model.document import Document


class BillingRun(Document):
	def before_save(self):
		if not self.currency:
			self.currency = "OMR"

	@frappe.whitelist()
	def get_storage_details(self):
		self.set("customer_details", [])
		if not self.period_from or not self.period_to:
			return
		
		# Fetch records from Storage Ledger between period_from and period_to
		records = frappe.db.get_all("Storage Leadger", 
			filters={
				"posting_date": ["between", [self.period_from, self.period_to]]
			},
			fields=["customer", "contract"]
		)
		
		seen = set()
		for record in records:
			if not record.customer or not record.contract:
				continue
			
			pair = (record.customer, record.contract)
			if pair not in seen:
				seen.add(pair)
				self.append("customer_details", {
					"customer": record.customer,
					"contract": record.contract,
					"select": 1
				})
				
	@frappe.whitelist()
	def get_selected_contract_details(self):
		self.set("billing_details", [])
		if not self.period_from or not self.period_to:
			return
		
		# Get selected customers and contracts
		selected_pairs = []
		if self.get("customer_details"):
			for row in self.get("customer_details"):
				if row.select:
					selected_pairs.append((row.customer, row.contract))
		
		if not selected_pairs:
			return
			
		for customer, contract in selected_pairs:
			records = frappe.db.get_all("Storage Leadger", 
				filters={
					"posting_date": ["between", [self.period_from, self.period_to]],
					"customer": customer,
					"contract": contract
				},
				fields=["customer", "warehouse", "reference_doctype", "reference_name", "item_code", "item_name", "pallet", "uom", "contract", "movement_type", "qty", "cbm_per_unit", "direction"]
			)
			
			for record in records:
				item_name = record.item_name
				if not item_name and record.item_code:
					item_name = frappe.db.get_value("Item", record.item_code, "item_name")
					
				self.append("billing_details", {
					"customer": record.customer,
					"warehouse": record.warehouse,
					"reference_doctype": record.reference_doctype,
					"reference_name": record.reference_name,
					"item_code": record.item_code,
					"item_name": item_name,
					"pallet": record.pallet,
					"uom": record.uom,
					"contract": record.contract,
					"movement_type": record.movement_type,
					"qty": record.qty,
					"cbm_per_unit": record.cbm_per_unit,
					"direction": record.direction,
					"check_rmeo": 1
				})
				
	def on_submit(self):
		self.create_sales_invoices()

	def create_sales_invoices(self):
		customer_lines = {}
		for line in self.billing_run_line:
			if not line.customer:
				continue
			if line.customer not in customer_lines:
				customer_lines[line.customer] = []
			customer_lines[line.customer].append(line)
			
		for customer, lines in customer_lines.items():
			si = frappe.new_doc("Sales Invoice")
			si.customer = customer
			si.posting_date = self.run_date or frappe.utils.today()
			si.due_date = frappe.utils.add_days(si.posting_date, 30)
			
			if si.meta.has_field("custom_billing_run"):
				si.custom_billing_run = self.name
				
			if si.meta.has_field("custom_doc_link_doctype_"):
				si.custom_doc_link_doctype_ = self.doctype
			if si.meta.has_field("custom_doc_link"):
				si.custom_doc_link = self.name
				
			for line in lines:
				item_code = line.item_code
				
				# Get HSN/SAC from Item
				gst_hsn_code = None
				if item_code:
					gst_hsn_code = frappe.db.get_value("Item", item_code, "gst_hsn_code")
				
				# Get Warehouse from Contract
				warehouse = None
				if line.contract:
					warehouse = frappe.db.get_value("Contract", line.contract, "custom_warehouse")
				
				desc = f"{line.charge_type} charge"
				actual_name = None
				if line.item_code:
					actual_name = frappe.db.get_value("Item", line.item_code, "item_name")
					desc += f" for Item {line.item_code}"
					if actual_name and actual_name != line.item_code:
						desc += f" ({actual_name})"
				if line.days:
					desc += f" ({line.days} days)"
				if line.direction:
					desc += f" [{line.direction}]"
					
				si.append("items", {
					"item_code": item_code,
					"item_name": actual_name or item_code,
					"description": desc,
					"qty": line.billed_qty or 1.0,
					"rate": line.rate or 0.0,
					"amount": line.billied_amount or 0.0,
					"warehouse": warehouse,
					"gst_hsn_code": gst_hsn_code
				})
				
			si.insert(ignore_permissions=True)
			si.submit()
			
			for line in lines:
				line.invoice = si.name
				line.db_update()
				
			for sum_line in self.billing_summerize_data:
				if sum_line.customer == customer:
					sum_line.invoice = si.name
					sum_line.db_update()
				
		self.db_set("status", "Invoiced")

	@frappe.whitelist()
	def calculate_bill(self):
		from frappe.utils import flt
		
		# Validation: Check if at least one row is checked
		selected_details = [row for row in self.get("billing_details") if row.check_rmeo]
		if not selected_details:
			frappe.throw(frappe._("Select at least one contract"))
			
		self.set("billing_run_line", [])
		
		contracts = {}
		for row in selected_details:
			if row.contract not in contracts:
				contracts[row.contract] = {"customer": row.customer, "rows": []}
			contracts[row.contract]["rows"].append(row)
				
		for contract, data in contracts.items():
			customer = data["customer"]
			contract_doc = frappe.get_doc("Contract", contract)
			
			if not contract_doc.get("custom_contract_tarrif"):
				continue
				
			for tariff in contract_doc.custom_contract_tarrif:
				if tariff.charge_type == "Storage":
					storage_lines = self.calculate_fifo_storage_qty_lines(
						customer=customer,
						contract=contract,
						period_from=self.period_from,
						period_to=self.period_to,
						billing_basis=(tariff.billing_basis or "").lower(),
						tariff_rate=tariff.rate or 0
					)
					for s_line in storage_lines:
						self.append("billing_run_line", {
							"contract": contract,
							"customer": customer,
							"charge_type": tariff.charge_type,
							"billing_basis": tariff.billing_basis,
							"direction": tariff.direction,
							"rate": tariff.rate or 0,
							"actual_qty": s_line["actual_qty"],
							"billed_qty": s_line["billed_qty"],
							"billied_amount": s_line["amount"],
							"is_one_time": tariff.is_one_time,
							"item_code": s_line["item_code"],
							"item_name": s_line.get("item_name") or s_line["item_code"],
							"days": s_line["days"]
						})
				else:
					basis = (tariff.billing_basis or "").lower()
					
					for row in data["rows"]:
						dir_val = (row.direction or "").lower()
						mov_val = (row.movement_type or "").lower()
						
						is_inbound = (dir_val == "inbound" or mov_val == "inbound")
						is_outbound = (dir_val == "outbound" or mov_val == "outbound")
						
						# Apply tariff direction filter
						if tariff.direction and tariff.direction != "Both":
							t_dir = tariff.direction.lower()
							if t_dir == "inbound" and not is_inbound:
								continue
							if t_dir == "outbound" and not is_outbound:
								continue
								
						# Calculate row qty
						if basis == "pallet":
							row_qty = flt(row.pallet or 0)
						elif basis == "cbm":
							row_qty = flt(row.cbm_per_unit or 0) * flt(row.qty or 0)
						else:
							row_qty = flt(row.qty or 0)
							
						if tariff.charge_type == "Fixed" or tariff.billing_basis == "Fixed":
							row_qty = 1.0
							
						rate = tariff.rate or 0
						amount = row_qty * rate
						
						item_code = row.item_code
						item_name = row.item_name
						if not item_name and item_code:
							item_name = frappe.db.get_value("Item", item_code, "item_name") or item_code
							
						self.append("billing_run_line", {
							"contract": contract,
							"customer": customer,
							"charge_type": tariff.charge_type,
							"billing_basis": tariff.billing_basis,
							"direction": "Outbound" if is_outbound else ("Inbound" if is_inbound else (row.direction or tariff.direction)),
							"rate": rate,
							"actual_qty": row_qty,
							"billed_qty": row_qty,
							"billied_amount": amount,
							"is_one_time": tariff.is_one_time,
							"item_code": item_code,
							"item_name": item_name,
							"days": ""
						})
						
					# Handle empty rows case for Fixed charges
					if not data["rows"] and (tariff.charge_type == "Fixed" or tariff.billing_basis == "Fixed"):
						rate = tariff.rate or 0
						self.append("billing_run_line", {
							"contract": contract,
							"customer": customer,
							"charge_type": tariff.charge_type,
							"billing_basis": tariff.billing_basis,
							"direction": tariff.direction or "Both",
							"rate": rate,
							"actual_qty": 1.0,
							"billed_qty": 1.0,
							"billied_amount": rate,
							"is_one_time": tariff.is_one_time,
							"item_code": "",
							"item_name": "",
							"days": ""
						})
							
		# Populate billing_summerize_data
		self.set("billing_summerize_data", [])
		summary_groups = {}
		for line in self.get("billing_run_line"):
			key = (line.contract, line.customer, line.charge_type, line.billing_basis, line.direction, line.rate)
			if key not in summary_groups:
				summary_groups[key] = {
					"contract": line.contract,
					"customer": line.customer,
					"charge_type": line.charge_type,
					"billing_basis": line.billing_basis,
					"direction": line.direction,
					"rate": line.rate,
					"actual_qty": 0.0,
					"billed_qty": 0.0,
					"billied_amount": 0.0,
					"is_one_time": line.is_one_time,
					"days_sum": 0.0
				}
				
			if line.days:
				try:
					summary_groups[key]["days_sum"] += flt(line.days)
				except Exception:
					pass
				
			summary_groups[key]["actual_qty"] += flt(line.actual_qty or 0)
			summary_groups[key]["billed_qty"] += flt(line.billed_qty or 0)
			summary_groups[key]["billied_amount"] += flt(line.billied_amount or 0)
			
		for key, s_data in summary_groups.items():
			days_val = ""
			if s_data["days_sum"] > 0:
				days_val = str(int(s_data["days_sum"])) if s_data["days_sum"].is_integer() else f"{s_data['days_sum']:.2f}"
				
			self.append("billing_summerize_data", {
				"contract": s_data["contract"],
				"customer": s_data["customer"],
				"charge_type": s_data["charge_type"],
				"billing_basis": s_data["billing_basis"],
				"direction": s_data["direction"],
				"rate": s_data["rate"],
				"actual_qty": s_data["actual_qty"],
				"billed_qty": s_data["billed_qty"],
				"billied_amount": s_data["billied_amount"],
				"is_one_time": s_data["is_one_time"],
				"days": days_val
			})
			
		total = 0
		for line in self.get("billing_run_line"):
			total += (line.billied_amount or 0)
			
		self.total_amount = total
		self.total_amt = total

	def calculate_fifo_storage_qty_lines(self, customer, contract, period_from, period_to, billing_basis, tariff_rate):
		from frappe.utils import date_diff, getdate, flt
		
		records = frappe.db.get_all("Storage Leadger",
			filters={
				"customer": customer,
				"contract": contract,
				"posting_date": ["<=", period_to]
			},
			fields=["posting_date", "item_code", "qty", "pallet", "cbm_per_unit", "direction"],
			order_by="posting_date asc, creation asc"
		)
		
		if not records:
			return []
			
		items_data = {}
		for r in records:
			item = r.item_code
			if item not in items_data:
				items_data[item] = {"inbounds": [], "outbounds": []}
			
			direction = (r.direction or "").lower()
			if direction in ["inbound", "inward"]:
				items_data[item]["inbounds"].append(r)
			elif direction in ["outbound", "outward"]:
				items_data[item]["outbounds"].append(r)
				
		total_days = date_diff(period_to, period_from) + 1
		if total_days <= 0:
			return []
			
		storage_lines = []
		
		for item_code, data in items_data.items():
			inbounds = data["inbounds"]
			outbounds = data["outbounds"]
			
			item_name = frappe.db.get_value("Item", item_code, "item_name") or item_code
			
			lots = []
			for ib in inbounds:
				qty = float(ib.qty or 0)
				pallet = float(ib.pallet or 0)
				cbm = float(ib.cbm_per_unit or 0) * qty
				
				lots.append({
					"posting_date": getdate(ib.posting_date),
					"qty": qty,
					"pallet": pallet,
					"cbm": cbm,
					"remaining_qty": qty,
					"remaining_pallet": pallet,
					"remaining_cbm": cbm
				})
				
			matches = []
			for ob in outbounds:
				ob_qty = float(ob.qty or 0)
				ob_date = getdate(ob.posting_date)
				
				qty_to_match = ob_qty
				for lot in lots:
					if qty_to_match <= 0:
						break
					if lot["remaining_qty"] <= 0:
						continue
					
					matched_qty = min(qty_to_match, lot["remaining_qty"])
					
					ratio = matched_qty / lot["qty"] if lot["qty"] > 0 else 0
					matched_pallet = ratio * lot["pallet"]
					matched_cbm = ratio * lot["cbm"]
					
					lot["remaining_qty"] -= matched_qty
					lot["remaining_pallet"] -= matched_pallet
					lot["remaining_cbm"] -= matched_cbm
					
					qty_to_match -= matched_qty
					
					stay_start = max(lot["posting_date"], getdate(period_from))
					stay_end = min(ob_date, getdate(period_to))
					
					if stay_end >= stay_start:
						days = date_diff(stay_end, stay_start) + 1
						
						if billing_basis == "pallet":
							basis_qty = matched_pallet
						elif billing_basis == "cbm":
							basis_qty = matched_cbm
						else:
							basis_qty = matched_qty
							
						matches.append({
							"days": days,
							"actual_qty": basis_qty
						})
						
			grouped_matches = {}
			for m in matches:
				days = m["days"]
				if days not in grouped_matches:
					grouped_matches[days] = 0.0
				grouped_matches[days] += m["actual_qty"]
				
			for days, act_qty in grouped_matches.items():
				billed_qty = flt((act_qty * days) / total_days, 4)
				amount = flt(billed_qty * tariff_rate, 2)
				storage_lines.append({
					"item_code": item_code,
					"item_name": item_name,
					"days": str(days),
					"actual_qty": act_qty,
					"billed_qty": billed_qty,
					"amount": amount
				})
				
		return storage_lines

	def calculate_fifo_storage_qty(self, customer, contract, period_from, period_to, billing_basis):
		from frappe.utils import date_diff, getdate
		
		records = frappe.db.get_all("Storage Leadger",
			filters={
				"customer": customer,
				"contract": contract,
				"posting_date": ["<=", period_to]
			},
			fields=["posting_date", "item_code", "qty", "pallet", "cbm_per_unit", "direction"],
			order_by="posting_date asc, creation asc"
		)
		
		if not records:
			return 0.0
			
		items_data = {}
		for r in records:
			item = r.item_code
			if item not in items_data:
				items_data[item] = {"inbounds": [], "outbounds": []}
			
			direction = (r.direction or "").lower()
			if direction in ["inbound", "inward"]:
				items_data[item]["inbounds"].append(r)
			elif direction in ["outbound", "outward"]:
				items_data[item]["outbounds"].append(r)
				
		total_days = date_diff(period_to, period_from) + 1
		if total_days <= 0:
			return 0.0
			
		total_basis_days = 0.0
		
		for item_code, data in items_data.items():
			inbounds = data["inbounds"]
			outbounds = data["outbounds"]
			
			lots = []
			for ib in inbounds:
				qty = float(ib.qty or 0)
				pallet = float(ib.pallet or 0)
				cbm = float(ib.cbm_per_unit or 0) * qty
				
				lots.append({
					"posting_date": getdate(ib.posting_date),
					"qty": qty,
					"pallet": pallet,
					"cbm": cbm,
					"remaining_qty": qty,
					"remaining_pallet": pallet,
					"remaining_cbm": cbm
				})
				
			for ob in outbounds:
				ob_qty = float(ob.qty or 0)
				ob_date = getdate(ob.posting_date)
				
				qty_to_match = ob_qty
				for lot in lots:
					if qty_to_match <= 0:
						break
					if lot["remaining_qty"] <= 0:
						continue
					
					matched_qty = min(qty_to_match, lot["remaining_qty"])
					
					ratio = matched_qty / lot["qty"] if lot["qty"] > 0 else 0
					matched_pallet = ratio * lot["pallet"]
					matched_cbm = ratio * lot["cbm"]
					
					lot["remaining_qty"] -= matched_qty
					lot["remaining_pallet"] -= matched_pallet
					lot["remaining_cbm"] -= matched_cbm
					
					qty_to_match -= matched_qty
					
					stay_start = max(lot["posting_date"], getdate(period_from))
					stay_end = min(ob_date, getdate(period_to))
					
					if stay_end >= stay_start:
						days = date_diff(stay_end, stay_start) + 1
						
						if billing_basis == "pallet":
							total_basis_days += matched_pallet * days
						elif billing_basis == "cbm":
							total_basis_days += matched_cbm * days
						else:
							total_basis_days += matched_qty * days
							
		return total_basis_days / total_days
	

@frappe.whitelist()
def export_billing_excel(name):
	doc = frappe.get_doc("Billing Run", name)
	
	import openpyxl
	from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
	from openpyxl.utils import get_column_letter
	import io
	
	wb = openpyxl.Workbook()
	
	# Sheet 1: WMSPro Detail Billing
	ws1 = wb.active
	ws1.title = "WMSPro Detail Billing"
	
	# Sheet 2: WmsPro Summerize Billing
	ws2 = wb.create_sheet(title="WmsPro Summerize Billing")
	
	# Sheet 3: WmsPro Customer Details
	ws3 = wb.create_sheet(title="WmsPro Customer Details")
	
	# Style helpers
	header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
	header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Navy Blue
	total_font = Font(name="Calibri", size=11, bold=True)
	center_align = Alignment(horizontal="center", vertical="center")
	left_align = Alignment(horizontal="left", vertical="center")
	right_align = Alignment(horizontal="right", vertical="center")
	
	thin_border = Border(
		left=Side(style='thin', color='D3D3D3'),
		right=Side(style='thin', color='D3D3D3'),
		top=Side(style='thin', color='D3D3D3'),
		bottom=Side(style='thin', color='D3D3D3')
	)
	
	double_bottom_border = Border(
		left=Side(style='thin', color='D3D3D3'),
		right=Side(style='thin', color='D3D3D3'),
		top=Side(style='thin', color='A0A0A0'),
		bottom=Side(style='double', color='000000')
	)
	
	# Populate Sheet 1
	headers1 = [
		"Contract", "Customer", "Item Code", "Item Name", "Charge Type",
		"Billing Basis", "Direction", "Days", "Actual Qty", "Billed Qty",
		"Rate", "Billed Amount", "Invoice"
	]
	
	for col_num, header in enumerate(headers1, 1):
		cell = ws1.cell(row=1, column=col_num)
		cell.value = header
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = center_align
		cell.border = thin_border
		
	for row_num, line in enumerate(doc.billing_run_line, 2):
		data = [
			line.contract,
			line.customer,
			line.item_code,
			line.item_name,
			line.charge_type,
			line.billing_basis,
			line.direction,
			line.days,
			line.actual_qty,
			line.billed_qty,
			line.rate,
			line.billied_amount,
			line.invoice
		]
		for col_num, val in enumerate(data, 1):
			cell = ws1.cell(row=row_num, column=col_num)
			cell.value = val
			cell.border = thin_border
			# Alignment
			if isinstance(val, (int, float)):
				cell.alignment = right_align
			else:
				cell.alignment = left_align
				
	# Total Row for Sheet 1
	last_row1 = len(doc.billing_run_line) + 1
	total_row1 = last_row1 + 1
	
	ws1.cell(row=total_row1, column=1, value="Total").font = total_font
	ws1.cell(row=total_row1, column=9, value=f"=SUM(I2:I{last_row1})").font = total_font
	ws1.cell(row=total_row1, column=10, value=f"=SUM(J2:J{last_row1})").font = total_font
	ws1.cell(row=total_row1, column=12, value=f"=SUM(L2:L{last_row1})").font = total_font
	
	for col_num in range(1, 14):
		cell = ws1.cell(row=total_row1, column=col_num)
		cell.border = double_bottom_border
		if col_num in [9, 10, 12]:
			cell.alignment = right_align
			
	# Auto-fit columns for Sheet 1
	for col in ws1.columns:
		max_len = max(len(str(cell.value or '')) for cell in col)
		col_letter = get_column_letter(col[0].column)
		ws1.column_dimensions[col_letter].width = max(max_len + 3, 10)
		
	# Populate Sheet 2
	headers2 = [
		"Contract", "Customer", "Charge Type", "Billing Basis", "Direction",
		"Days", "Actual Qty", "Billed Qty", "Rate", "Billed Amount", "Invoice"
	]
	
	for col_num, header in enumerate(headers2, 1):
		cell = ws2.cell(row=1, column=col_num)
		cell.value = header
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = center_align
		cell.border = thin_border
		
	for row_num, line in enumerate(doc.billing_summerize_data, 2):
		data = [
			line.contract,
			line.customer,
			line.charge_type,
			line.billing_basis,
			line.direction,
			line.days,
			line.actual_qty,
			line.billed_qty,
			line.rate,
			line.billied_amount,
			line.invoice
		]
		for col_num, val in enumerate(data, 1):
			cell = ws2.cell(row=row_num, column=col_num)
			cell.value = val
			cell.border = thin_border
			if isinstance(val, (int, float)):
				cell.alignment = right_align
			else:
				cell.alignment = left_align
				
	# Total Row for Sheet 2
	last_row2 = len(doc.billing_summerize_data) + 1
	total_row2 = last_row2 + 1
	
	ws2.cell(row=total_row2, column=1, value="Total").font = total_font
	ws2.cell(row=total_row2, column=7, value=f"=SUM(G2:G{last_row2})").font = total_font
	ws2.cell(row=total_row2, column=8, value=f"=SUM(H2:H{last_row2})").font = total_font
	ws2.cell(row=total_row2, column=10, value=f"=SUM(J2:J{last_row2})").font = total_font
	
	for col_num in range(1, 12):
		cell = ws2.cell(row=total_row2, column=col_num)
		cell.border = double_bottom_border
		if col_num in [7, 8, 10]:
			cell.alignment = right_align
			
	# Auto-fit columns for Sheet 2
	for col in ws2.columns:
		max_len = max(len(str(cell.value or '')) for cell in col)
		col_letter = get_column_letter(col[0].column)
		ws2.column_dimensions[col_letter].width = max(max_len + 3, 10)
		
	# Populate Sheet 3 (WmsPro Customer Details)
	customer_summary = {}
	for line in doc.billing_run_line:
		key = (line.customer, line.charge_type, line.billing_basis, line.direction, line.rate)
		if key not in customer_summary:
			customer_summary[key] = {
				"customer": line.customer,
				"charge_type": line.charge_type,
				"billing_basis": line.billing_basis,
				"direction": line.direction,
				"rate": line.rate,
				"actual_qty": 0.0,
				"billed_qty": 0.0,
				"amount": 0.0,
				"invoices": set()
			}
		customer_summary[key]["actual_qty"] += float(line.actual_qty or 0)
		customer_summary[key]["billed_qty"] += float(line.billed_qty or 0)
		customer_summary[key]["amount"] += float(line.billied_amount or 0)
		if line.invoice:
			customer_summary[key]["invoices"].add(line.invoice)
			
	headers3 = [
		"Contract", "Customer", "Charge Type", "Billing Basis", "Direction",
		"Days", "Actual Qty", "Billed Qty", "Rate", "Billed Amount", "Invoice"
	]
	
	for col_num, header in enumerate(headers3, 1):
		cell = ws3.cell(row=1, column=col_num)
		cell.value = header
		cell.font = header_font
		cell.fill = header_fill
		cell.alignment = center_align
		cell.border = thin_border
		
	row_num3 = 2
	for key, s_data in customer_summary.items():
		inv_str = ", ".join(sorted(list(s_data["invoices"])))
		data = [
			"", # Contract combined
			s_data["customer"],
			s_data["charge_type"],
			s_data["billing_basis"],
			s_data["direction"],
			"", # Days combined
			s_data["actual_qty"],
			s_data["billed_qty"],
			s_data["rate"],
			s_data["amount"],
			inv_str
		]
		for col_num, val in enumerate(data, 1):
			cell = ws3.cell(row=row_num3, column=col_num)
			cell.value = val
			cell.border = thin_border
			if isinstance(val, (int, float)):
				cell.alignment = right_align
			else:
				cell.alignment = left_align
				
		row_num3 += 1
		
	# Total Row for Sheet 3
	last_row3 = row_num3 - 1
	total_row3 = row_num3
	
	ws3.cell(row=total_row3, column=1, value="Total").font = total_font
	ws3.cell(row=total_row3, column=7, value=f"=SUM(G2:G{last_row3})").font = total_font
	ws3.cell(row=total_row3, column=8, value=f"=SUM(H2:H{last_row3})").font = total_font
	ws3.cell(row=total_row3, column=10, value=f"=SUM(J2:J{last_row3})").font = total_font
	
	for col_num in range(1, 12):
		cell = ws3.cell(row=total_row3, column=col_num)
		cell.border = double_bottom_border
		if col_num in [7, 8, 10]:
			cell.alignment = right_align
			
	# Auto-fit columns for Sheet 3
	for col in ws3.columns:
		max_len = max(len(str(cell.value or '')) for cell in col)
		col_letter = get_column_letter(col[0].column)
		ws3.column_dimensions[col_letter].width = max(max_len + 3, 10)
		
	# Save to buffer
	file_stream = io.BytesIO()
	wb.save(file_stream)
	file_stream.seek(0)
	
	# Set Response headers to trigger download in Frappe
	frappe.response['filename'] = f"WMSPro_Billing_Run_{doc.name}.xlsx"
	frappe.response['filecontent'] = file_stream.getvalue()
	frappe.response['type'] = 'binary'


	
